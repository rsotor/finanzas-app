# -*- coding: utf-8 -*-
import sys, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)
ws = wb.create_sheet('Léeme', 0)
AZ = 'FF1F3864'
L = [
 ('t', 'FINANZAS FAMILIARES — plantilla con datos de ejemplo'),
 ('s', 'Todas las cifras son INVENTADAS (una familia ficticia: Ana, Luis y su hijo Leo). Sustitúyelas por las tuyas.'),
 ('', ''),
 ('h', 'QUÉ HAY DENTRO'),
 ('p', 'Dos partes que se alimentan solas:'),
 ('p', '  • FINANZAS (lo que tienes hoy): Patrimonio Neto, Ingresos - Gastos, Presupuesto, Gráficos, Fondo de emergencia.'),
 ('p', '  • PLAN (a dónde quieres llegar): Panel, Objetivos, Proyección. Lee tu ahorro y lo ya invertido de la parte de finanzas.'),
 ('', ''),
 ('h', 'EN QUÉ ORDEN RELLENARLO'),
 ('p', '1. Patrimonio Neto — lo que tienes y lo que debes. Pon la FECHA de hoy en E1: los plazos de los objetivos se cuentan desde ese año.'),
 ('p', '2. Ingresos - Gastos — cada concepto en la columna B (si es mensual) o en la C (si es anual), nunca en las dos. La G es para tus notas.'),
 ('p', '     Truco: saca los importes del extracto del banco de los últimos 12 meses; los recibos anuales o estacionales (IBI, seguros, gas) mejor en anual.'),
 ('p', '3. Presupuesto — la columna D (en azul) es lo que QUIERES gastar. El resto se copia solo de «Ingresos - Gastos».'),
 ('p', '4. Fondo de emergencia — cuánto aguantaríais sin ingresos. Rellena el paro estimado y los recortes que harías.'),
 ('p', '5. Panel — rentabilidad supuesta y aportación mensual de cada cartera, y los supuestos (inflación, impuestos…). Es lo único que tocas para "jugar".'),
 ('p', '6. Objetivos — cada meta con su año, importe en euros de HOY y la cartera de la que sale.'),
 ('p', '7. Proyección — no se edita: año a año, cuánto hay en cada cartera y si llega para cada objetivo.'),
 ('', ''),
 ('h', 'REGLAS PARA NO ROMPERLO'),
 ('p', '• Edita solo las celdas con FONDO VERDE (Panel y Objetivos), la columna D del Presupuesto y los importes de las demás hojas.'),
 ('p', '• NO insertes ni borres filas: cada grupo trae filas vacías de reserva ya incluidas en las sumas. Si te sobra una, déjala en blanco.'),
 ('p', '• Elige una herramienta (Excel o Google Sheets) y quédate con ella: pasar el fichero de una a otra suele estropear los gráficos.'),
 ('p', '• En Google Sheets: súbelo a Drive y ábrelo con «Abrir con → Hojas de cálculo de Google».'),
 ('', ''),
 ('h', 'CONCEPTOS CLAVE'),
 ('p', '• Carteras por plazo ("cubos"): el dinero que necesitas en 5 años no puede ir donde el de 30. Lo cercano va en algo estable (cartera conservadora);'),
 ('p', '   lo lejano, en renta variable (indexada global), que rinde más a largo plazo pero puede caer un 20-30% en un mal año.'),
 ('p', '• Euros de hoy: los objetivos se escriben con precios actuales; el libro les suma la inflación hasta su año.'),
 ('p', '• Rentabilidad mínima para no perder = inflación ÷ (1 − impuestos). Por debajo de eso, tu dinero pierde poder de compra.'),
 ('p', '• Tipos de objetivo: «Único» se paga de golpe en su año; «Renta» es una cantidad al mes durante X años; «Jubilación» es una renta que además'),
 ('p', '   marca el año en que dejas de aportar.'),
 ('p', '• Fondo de emergencia: dinero líquido y separado para imprevistos o para quedarse sin ingresos. No se invierte en bolsa.'),
 ('', ''),
 ('h', 'LÍMITES (léelo antes de fiarte de los números)'),
 ('p', '• NO es asesoramiento financiero. Es una calculadora: los resultados dependen de los supuestos que pongas.'),
 ('p', '• Las rentabilidades son SUPUESTOS para jugar, no promesas. Rentabilidades pasadas no garantizan las futuras.'),
 ('p', '• Impuestos simplificados: se aplica un tipo medio a la plusvalía al vender. En España la base del ahorro va por tramos (del 19% al 30%)'),
 ('p', '   y Hacienda calcula la plusvalía por FIFO (lo primero que compraste es lo primero que vendes); el libro usa el coste medio como aproximación.'),
 ('p', '• «Necesita €/mes» es conservador: no descuenta lo que ya tienes invertido ni la subida anual de la aportación. Si sale justo, probablemente vas bien.'),
 ('p', '• No incluye la pensión pública: el objetivo de jubilación es el COMPLEMENTO que quieres añadirle.'),
]
for i, (k, t) in enumerate(L, start=1):
    c = ws.cell(i, 1, t)
    if k == 't': c.font = Font(bold=True, size=16, color=AZ)
    elif k == 's': c.font = Font(italic=True, size=10, color='FFC00000')
    elif k == 'h':
        c.font = Font(bold=True, size=11, color='FFFFFFFF')
        c.fill = PatternFill('solid', fgColor='FF4472C4')
    else: c.font = Font(size=10)
ws.column_dimensions['A'].width = 150
ws.sheet_view.showGridLines = False
wb.active = 0
for w in wb.worksheets: w.sheet_view.tabSelected = (w.title == 'Léeme')
wb.save(OUT)
print('ok', wb.sheetnames)
