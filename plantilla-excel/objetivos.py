# -*- coding: utf-8 -*-
"""Anade las pestanas «Objetivos» y «Proyeccion» a Finanzas personales.xlsx.

   MODELO POR CARTERAS (cubos): cada objetivo se asigna a una cartera con su
   propia rentabilidad, y la proyeccion lleva una bolsa separada por cartera.
   El dinero del coche a 5 anos NO compone al mismo ritmo que la jubilacion.

   TIPOS de objetivo:
     Unico      -> se saca todo de golpe en su ano
     Renta      -> cuota anual (indexada a inflacion) durante N anos
     Jubilacion -> igual que Renta, y ademas marca el ano en que dejas de aportar

   Los objetivos se guardan en ANO ABSOLUTO; el «dentro de» se calcula desde
   «Patrimonio Neto»!E1, asi la tabla no caduca.

   Fiscalidad: se lleva la BASE (coste de adquisicion) de cada bolsa por el
   metodo de coste medio; al rescatar, la parte de plusvalia tributa al tipo
   indicado.

   Escribe formulas Y calcula lo mismo en Python para inyectar valores cacheados
   y verificar el modelo."""
import sys, json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, Reference

SRC, OUT, VALJSON, CALC, REFS, NOMBRE_ESC = sys.argv[1:7]
REF = json.load(open(REFS))
sys.path.insert(0, __import__('os').path.dirname(__import__('os').path.abspath(__file__)))
from escenarios import ESCENARIOS
ESC = ESCENARIOS[NOMBRE_ESC]
INF, IMP, INCR = ESC['INF'], ESC['IMP'], ESC['INCR']

AZUL_OSC, AZUL, GRIS, NARANJA, VERDE = 'FF1F3864', 'FF4472C4', 'FFF2F2F2', 'FFFCE4D6', 'FFE2EFDA'
EUR = '#,##0.00\\ "€"'
EUR0 = '#,##0\\ "€"'
PCT = '0.00%'
EDITABLE = PatternFill('solid', fgColor='FFDDF3E0')
REFF = PatternFill('solid', fgColor='FFF7F7F7')

CELDA_PP_ANUAL = REF['pp_anual']
CELDA_AHORRO = REF['ahorro']


def titulo(ws, ref, txt):
    ws[ref] = txt
    ws[ref].font = Font(bold=True, size=15, color=AZUL_OSC)


def nota(ws, ref, txt):
    ws[ref] = txt
    ws[ref].font = Font(size=9, color='FF808080')


def banda(ws, row, cols, txt):
    for c in cols:
        ws[f'{c}{row}'].fill = PatternFill('solid', fgColor=AZUL)
        ws[f'{c}{row}'].font = Font(bold=True, size=12, color='FFFFFFFF')
    ws[f'{cols[0]}{row}'] = txt


def cabecera(ws, row, pares, alto=34):
    for col, txt in pares:
        c = ws[f'{col}{row}']
        c.value = txt
        c.fill = PatternFill('solid', fgColor=AZUL_OSC)
        c.font = Font(bold=True, size=10, color='FFFFFFFF')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = alto


def editable(c):
    c.fill = EDITABLE
    c.font = Font(bold=True, size=10, color='FF0070C0')
    c.border = Border(*[Side(style='thin', color='FFAAAAAA')] * 4)


wb = openpyxl.load_workbook(SRC)
for h in ('Panel', 'Objetivos', 'Proyección'):
    if h in wb.sheetnames:
        del wb[h]

# ---- carteras: (nombre, rentabilidad, ya_invertido_formula, extra_ini, mens, nota)
CAR0, CARN = 19, 21
_IDS = {c[0]: c[1] for c in ESC['CARTERAS']}
CARTERAS = [(nom, rent, (REF['inv'][inv] if inv else None), extra, mens, nt)
            for (_id, nom, rent, inv, extra, mens, nt) in ESC['CARTERAS']]
NCAR = len(CARTERAS)

OBJ0, OBJN = 26, 45
TOT = OBJN + 2

# =========================================================================
#  HOJA 1 - OBJETIVOS
# =========================================================================
ob = wb.create_sheet('Objetivos')
titulo(ob, 'A1', 'OBJETIVOS')
nota(ob, 'A2', 'Rellena solo las celdas con FONDO VERDE. Importes en euros de HOY y fechas en año real. '
               'Cada objetivo va asignado a una CARTERA, y cada cartera compone a su propio ritmo: '
               'el dinero a 5 años no puede ir donde el de 30.')

banda(ob, 4, list('ABC'), 'PARÁMETROS')
PARAMS = [
    (5,  'Año de partida', None, '0', 'calc',
     'Sale de la fecha de «Patrimonio Neto». Es el año 0: los plazos se cuentan desde aquí'),
    (6,  'Inflación anual', None, PCT, 'calc', 'Se toca en la pestaña «Panel»'),
    (7,  'Impuestos sobre la plusvalía', None, PCT, 'calc', 'Se toca en la pestaña «Panel»'),
    (8,  'Rentabilidad MÍNIMA para no perder', None, PCT, 'calc',
     'Inflación ÷ (1 − impuestos). Es el SUELO: por debajo pierdes poder adquisitivo'),
    (9,  'Incremento anual de la aportación', None, PCT, 'calc', 'Se toca en la pestaña «Panel»'),
    (10, 'Dejo de aportar dentro de (años)', None, '0', 'calc', 'Se toca en la pestaña «Panel»'),
    (11, '→ Plazo que se aplica (años)', None, '0', 'calc', 'El que usa la proyección'),
    (12, '→ …es decir, el año', None, '0', 'calc', 'El año en que dejas de aportar'),
    (13, 'Tu ahorro real (de «Ingresos - Gastos»)', None, EUR, 'calc',
     'Lo que de verdad te sobra cada mes según tus cuentas'),
    (14, '→ Aportas al mes (todas las carteras)', None, EUR, 'calc', 'Suma de la columna «aportación mensual»'),
    (15, '→ % de tu ahorro que inviertes', None, PCT, 'calc',
     'Sin el plan de pensiones, que ya está contado como gasto'),
]
for row, etiq, val, fmt, modo, ayuda in PARAMS:
    ob[f'A{row}'] = etiq
    ob[f'A{row}'].font = Font(bold=True, size=10, color=AZUL_OSC)
    c = ob[f'B{row}']
    c.number_format = fmt
    if modo == 'edit':
        if val is not None:
            c.value = val
        editable(c)
    else:
        c.font = Font(bold=True, size=11, color=AZUL_OSC)
    ob[f'C{row}'] = ayuda
    ob[f'C{row}'].font = Font(size=9, color='FF808080')

ob['B5'] = "=YEAR('Patrimonio Neto'!E1)"
ob['B6'] = '=Panel!B17'
ob['B7'] = '=Panel!B18'
ob['B8'] = '=B6/(1-B7)'
ob['B9'] = '=Panel!B19'
ob['B10'] = '=IF(Panel!B20="","",Panel!B20)'
ob['B11'] = f'=IF(B10="",IF(MIN(N{OBJ0}:N{OBJN})=9999,999,MIN(N{OBJ0}:N{OBJN})),B10)'
ob['B12'] = '=B5+B11'
ob['B13'] = f'={CELDA_AHORRO}'
ob['B14'] = f'=SUM(E{CAR0}:E{CARN})'
ob['B15'] = f'=IF(B13=0,0,(B14-E{CAR0 + 2})/B13)'

# ---- bloque CARTERAS -----------------------------------------------------
banda(ob, 17, list('ABCDEFGH'), 'TUS CARTERAS  (cada una compone a su ritmo)')
cabecera(ob, 18, [('A', 'CARTERA'), ('B', 'RENTABILIDAD'), ('C', 'YA INVERTIDO\n(€)'),
                  ('D', 'APORTAS AL\nEMPEZAR (€)'), ('E', 'APORTACIÓN\nMENSUAL (€)'),
                  ('F', '→ NECESARIO\nAL MES (€)'), ('G', 'ESTADO'), ('H', 'NOTAS')])
for i, (nom, rent, inv, extra, mens, nt) in enumerate(CARTERAS):
    r = CAR0 + i
    ob[f'A{r}'] = nom
    ob[f'A{r}'].font = Font(bold=True, size=10, color=AZUL_OSC)
    ob[f'B{r}'] = f'=Panel!B{9 + i}'
    ob[f'B{r}'].number_format = PCT
    ob[f'B{r}'].font = Font(bold=True, size=10, color=AZUL_OSC)
    if inv:
        ob[f'C{r}'] = f'={inv}'
        ob[f'C{r}'].font = Font(bold=True, size=10, color=AZUL_OSC)
    else:
        ob[f'C{r}'] = 0
        editable(ob[f'C{r}'])
    ob[f'C{r}'].number_format = EUR
    ob[f'D{r}'] = extra
    ob[f'D{r}'].number_format = EUR0
    editable(ob[f'D{r}'])
    ob[f'E{r}'] = f'=Panel!C{9 + i}'
    ob[f'E{r}'].font = Font(bold=True, size=10, color=AZUL_OSC)
    ob[f'E{r}'].number_format = EUR
    ob[f'F{r}'] = f'=SUMIFS($L${OBJ0}:$L${OBJN},$B${OBJ0}:$B${OBJN},$A{r})'
    ob[f'F{r}'].number_format = EUR
    ob[f'F{r}'].font = Font(size=10, color=AZUL_OSC)
    ob[f'G{r}'] = (f'=IF(F{r}=0,"— sin objetivos",IF(E{r}>=F{r},"✓ cubierta",'
                   f'"⚠ faltan "&TEXT(F{r}-E{r},"#,##0")&" €/mes"))')
    ob[f'G{r}'].font = Font(bold=True, size=10, color='FFC00000')
    ob[f'H{r}'] = nt
    ob[f'H{r}'].font = Font(size=9, color='FF808080')
    ob[f'H{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    ob.row_dimensions[r].height = 30
rT = CARN + 1
ob[f'A{rT}'] = 'TOTAL'
for col in 'ABCDEFGH':
    ob[f'{col}{rT}'].fill = PatternFill('solid', fgColor=NARANJA)
    ob[f'{col}{rT}'].font = Font(bold=True, size=11, color=AZUL_OSC)
for col in 'CDEF':
    ob[f'{col}{rT}'] = f'=SUM({col}{CAR0}:{col}{CARN})'
    ob[f'{col}{rT}'].number_format = EUR

dvc = DataValidation(type='list', formula1=f'=$A${CAR0}:$A${CARN}', allow_blank=True,
                     prompt='A qué cartera va este objetivo. Cada una compone a su rentabilidad.',
                     promptTitle='Cartera')
ob.add_data_validation(dvc)

# ---- bloque OBJETIVOS ----------------------------------------------------
banda(ob, 24, list('ABCDEFGHIJKLM'), 'TUS OBJETIVOS')
cabecera(ob, 25, [('A', 'OBJETIVO'), ('B', 'CARTERA'), ('C', 'TIPO'), ('D', 'AÑO'), ('E', 'DENTRO\nDE'),
                  ('F', 'IMPORTE HOY\n(€)'), ('G', 'RENTA MENSUAL\nHOY (€)'), ('H', 'AÑOS DE\nCONSUMO'),
                  ('I', 'TOTAL A RETIRAR\n(€ de su año)'), ('J', 'CAPITAL NECESARIO\nAL LLEGAR (€)'),
                  ('K', 'SI LO PAGAS HOY\nDE GOLPE (€)'), ('L', 'APORTACIÓN MENSUAL\nNECESARIA (€)'),
                  ('M', 'NOTAS'), ('N', 'aux:\njubilación')])

#   (nombre, cartera, tipo, ano, importe_hoy, renta_mes, anios_consumo, nota)
OBJ = [(nom, _IDS.get(car), tipo, anio, imp, renta, cons, nt) for (nom, car, tipo, anio, imp, renta, cons, nt) in ESC['OBJ']]
for i, (nom, car, tipo, anio, imp, renta, cons, nt) in enumerate(OBJ):
    r = OBJ0 + i
    ob[f'A{r}'], ob[f'B{r}'], ob[f'C{r}'], ob[f'D{r}'] = nom, car, tipo, anio
    ob[f'F{r}'], ob[f'G{r}'], ob[f'H{r}'] = imp, renta, cons
    ob[f'M{r}'] = nt

RENTOBJ = f'IFERROR(VLOOKUP($B{{r}},$A${CAR0}:$B${CARN},2,FALSE),$B$8)'
for r in range(OBJ0, OBJN + 1):
    for col in 'ABCDFGH':
        ob[f'{col}{r}'].fill = EDITABLE
        ob[f'{col}{r}'].border = Border(*[Side(style='thin', color='FFCCCCCC')] * 4)
    for col, fmt in (('F', EUR0), ('G', EUR0), ('I', EUR), ('J', EUR), ('K', EUR), ('L', EUR)):
        ob[f'{col}{r}'].number_format = fmt
    for col in 'DH':
        ob[f'{col}{r}'].number_format = '0'
    ob[f'A{r}'].font = Font(size=10)
    ob[f'B{r}'].font = Font(size=9)
    ob[f'M{r}'].font = Font(size=9, color='FFC00000')
    ob[f'M{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    for col in 'EIJKL':
        ob[f'{col}{r}'].font = Font(size=10, color=AZUL_OSC)
    ob[f'J{r}'].font = Font(bold=True, size=10, color=AZUL_OSC)
    ro = RENTOBJ.format(r=r)
    esr = f'OR($C{r}="Renta",$C{r}="Jubilación")'
    # OJO: 0 y no "" — la Proyección hace E+H dentro de un SUMPRODUCT
    ob[f'E{r}'] = f'=IF($C{r}="",0,D{r}-$B$5)'
    ob[f'E{r}'].number_format = '0;-0;;'
    ob[f'I{r}'] = (f'=IF($C{r}="","",IF({esr},'
                   f'G{r}*12*(1+$B$6)^E{r}*((1+$B$6)^H{r}-1)/$B$6,'
                   f'F{r}*(1+$B$6)^E{r}))')
    ob[f'J{r}'] = (f'=IF($C{r}="","",IF({esr},'
                   f'G{r}*12*(1+$B$6)^E{r}*IF(ABS($B$6-{ro})<0.0000001,H{r},'
                   f'(1-((1+$B$6)/(1+{ro}))^H{r})/(1-(1+$B$6)/(1+{ro}))),I{r}))')
    ob[f'K{r}'] = f'=IF($C{r}="","",J{r}/(1+{ro})^E{r})'
    ob[f'L{r}'] = f'=IF($C{r}="","",IF(E{r}<=0,"",((J{r}*{ro})/((1+{ro})^E{r}-1))/12))'
    ob[f'N{r}'] = f'=IF($C{r}="Jubilación",E{r},9999)'
    ob[f'N{r}'].font = Font(size=9, color='FFBBBBBB')
    ob[f'N{r}'].number_format = '0'
    dvc.add(f'B{r}')

ob[f'A{TOT}'] = 'TOTAL'
for col in 'ABCDEFGHIJKLMN':
    ob[f'{col}{TOT}'].fill = PatternFill('solid', fgColor=NARANJA)
    ob[f'{col}{TOT}'].font = Font(bold=True, size=12, color=AZUL_OSC)
for col in 'IJKL':
    ob[f'{col}{TOT}'] = f'=SUM({col}{OBJ0}:{col}{OBJN})'
    ob[f'{col}{TOT}'].number_format = EUR

dvt = DataValidation(type='list', formula1='"Único,Renta,Jubilación"', allow_blank=True,
                     prompt='Único = de golpe en su año.\nRenta = una cuota al año durante X años, y deja de '
                            'descontar al terminar.\nJubilación = igual que Renta, y además marca el año en que '
                            'dejas de aportar.',
                     promptTitle='Tipo de objetivo')
ob.add_data_validation(dvt)
dvt.add(f'C{OBJ0}:C{OBJN}')

for col, w in (('A', 30), ('B', 18), ('C', 12), ('D', 9), ('E', 9), ('F', 14), ('G', 15), ('H', 11),
               ('I', 17), ('J', 18), ('K', 17), ('L', 19), ('M', 55), ('N', 10)):
    ob.column_dimensions[col].width = w
ob.freeze_panes = f'A{OBJ0}'

# =========================================================================
#  HOJA 2 - PROYECCION  (una bolsa por cartera)
# =========================================================================
pr = wb.create_sheet('Proyección')
ANIOS = 85
R0 = 6
RN = R0 + ANIOS - 1
titulo(pr, 'A1', 'PROYECCIÓN POR CARTERAS')
nota(pr, 'A2', 'NO se edita nada aquí. Cada cartera es una bolsa independiente que compone a su propia '
               'rentabilidad; los objetivos rescatan de la suya. A la izquierda el total, a la derecha el detalle.')
SCC = [openpyxl.utils.get_column_letter(7 + k * 6 + 5) for k in range(NCAR)]
SUMSC = '+'.join(f'SUM({c}{R0}:{c}{RN})' for c in SCC)
PRIMER = ('MIN(' + ','.join(f'SUMPRODUCT(MIN(({c}{R0}:{c}{RN}>0.01)*A{R0}:A{RN}+({c}{R0}:{c}{RN}<=0.01)*999999))' for c in SCC) + ')')
pr['A3'] = (f'=IF({SUMSC}<0.01,"✓ Todos los objetivos se pagan",'
            f'"⚠ NO LLEGAS: el primer objetivo que se queda sin pagar es en {{}}"&{PRIMER})'.replace('{}',''))
pr['A3'].font = Font(bold=True, size=12, color='FFC00000')

RES = [('A', 'AÑO'), ('B', 'DENTRO\nDE'), ('C', 'TOTAL\nACUMULADO'), ('D', 'TOTAL NETO\n(si vendes)'),
       ('E', 'APORTAS\nEL AÑO'), ('F', 'RETIRAS\nEL AÑO')]
cabecera(pr, 5, RES, alto=30)
pr['A4'] = 'RESUMEN'
pr['A4'].font = Font(bold=True, size=10, color=AZUL_OSC)
for i, (nom, *_ ) in enumerate(CARTERAS):
    base = 7 + i * 6
    letras = [openpyxl.utils.get_column_letter(base + k) for k in range(6)]
    pr[f'{letras[0]}4'] = nom
    pr[f'{letras[0]}4'].font = Font(bold=True, size=11, color='FFFFFFFF')
    for L in letras:
        pr[f'{L}4'].fill = PatternFill('solid', fgColor=AZUL)
    cabecera(pr, 5, list(zip(letras, ['APORTAS', 'VALOR', 'BASE\n(lo aportado)',
                                      'RETIRAS', 'COSTE REAL\n(con impuesto)',
                                      '⚠ SIN\nCUBRIR'])), alto=30)

APORT = ('IF(B{r}>=Objetivos!$B$11,0,'
         'IF(B{r}=1,Objetivos!$E${cr}*12+Objetivos!$C${cr}+Objetivos!$D${cr},'
         'Objetivos!$E${cr}*12*(1+Objetivos!$B$9)^(B{r}-1)))')


def unicos(cr):
    return (f'SUMIFS(Objetivos!$I${OBJ0}:$I${OBJN},Objetivos!$E${OBJ0}:$E${OBJN},B{{r}},'
            f'Objetivos!$C${OBJ0}:$C${OBJN},"Único",Objetivos!$B${OBJ0}:$B${OBJN},Objetivos!$A${cr})')


def rentas(cr):
    return (f'SUMPRODUCT(((Objetivos!$C${OBJ0}:$C${OBJN}="Renta")+'
            f'(Objetivos!$C${OBJ0}:$C${OBJN}="Jubilación"))*'
            f'(Objetivos!$B${OBJ0}:$B${OBJN}=Objetivos!$A${cr})*'
            f'(B{{r}}>=Objetivos!$E${OBJ0}:$E${OBJN})*'
            f'(B{{r}}<Objetivos!$E${OBJ0}:$E${OBJN}+Objetivos!$H${OBJ0}:$H${OBJN})*'
            f'Objetivos!$G${OBJ0}:$G${OBJN})*12*(1+Objetivos!$B$6)^B{{r}}')


for i in range(ANIOS):
    r = R0 + i
    a = i + 1
    p = r - 1
    pr[f'A{r}'] = f'=Objetivos!$B$5+B{r}'
    pr[f'B{r}'] = a
    for k in range(NCAR):
        cr = CAR0 + k
        base = 7 + k * 6
        AP, VA, BA, RE, CO, SC = [openpyxl.utils.get_column_letter(base + j) for j in range(6)]
        pedido = unicos(cr).format(r=r) + '+' + rentas(cr).format(r=r)
        pr[f'{AP}{r}'] = '=' + APORT.format(r=r, cr=cr)
        pr[f'{VA}{r}'] = (f'={AP}{r}' if a == 1 else
                          f'=MAX(0,({VA}{p}-{CO}{p})*(1+Objetivos!$B${cr})+{AP}{r})')
        # base = coste de adquisicion, metodo de coste medio
        pr[f'{BA}{r}'] = (f'={AP}{r}' if a == 1 else
                          f'=MAX(0,{BA}{p}-IF({VA}{p}=0,0,{RE}{p}*{BA}{p}/{VA}{p})+{AP}{r})')
        pr[f'{RE}{r}'] = f'=MIN({VA}{r},' + pedido + ')'
        pr[f'{SC}{r}'] = f'=MAX(0,' + pedido + f'-{VA}{r})'
        pr[f'{CO}{r}'] = (f'={RE}{r}+IF({VA}{r}=0,0,{RE}{r}*MAX(0,{VA}{r}-{BA}{r})/{VA}{r})*Objetivos!$B$7')
        for L in (AP, VA, BA, RE, CO, SC):
            pr[f'{L}{r}'].number_format = EUR
            pr[f'{L}{r}'].font = Font(size=9)
        pr[f'{VA}{r}'].font = Font(bold=True, size=9, color=AZUL_OSC)
        pr[f'{RE}{r}'].font = Font(size=9, color='FFC00000')
        pr[f'{SC}{r}'].font = Font(bold=True, size=9, color='FFC00000')
    vals = [openpyxl.utils.get_column_letter(7 + k * 6 + 1) for k in range(NCAR)]
    bases = [openpyxl.utils.get_column_letter(7 + k * 6 + 2) for k in range(NCAR)]
    aps = [openpyxl.utils.get_column_letter(7 + k * 6) for k in range(NCAR)]
    res = [openpyxl.utils.get_column_letter(7 + k * 6 + 3) for k in range(NCAR)]
    pr[f'C{r}'] = '=' + '+'.join(f'{v}{r}' for v in vals)
    pr[f'D{r}'] = ('=' + '+'.join(f'({v}{r}-MAX(0,{v}{r}-{b}{r})*Objetivos!$B$7)'
                                  for v, b in zip(vals, bases)))
    pr[f'E{r}'] = '=' + '+'.join(f'{x}{r}' for x in aps)
    pr[f'F{r}'] = '=' + '+'.join(f'{x}{r}' for x in res)
    pr[f'A{r}'].number_format = '0'
    pr[f'A{r}'].font = Font(bold=True, size=10, color=AZUL_OSC)
    pr[f'B{r}'].number_format = '0'
    pr[f'B{r}'].font = Font(size=9, color='FF999999')
    for col in 'CDEF':
        pr[f'{col}{r}'].number_format = EUR
        pr[f'{col}{r}'].font = Font(size=10)
    for col in 'CD':
        pr[f'{col}{r}'].font = Font(bold=True, size=10, color=AZUL_OSC)
    pr[f'F{r}'].font = Font(bold=True, size=10, color='FFC00000')
    if i % 2:
        for c in range(1, 7 + NCAR * 6):
            pr.cell(r, c).fill = PatternFill('solid', fgColor=GRIS)

for col, w in [('A', 8), ('B', 8), ('C', 15), ('D', 15), ('E', 13), ('F', 13)]:
    pr.column_dimensions[col].width = w
for k in range(NCAR * 6):
    pr.column_dimensions[openpyxl.utils.get_column_letter(7 + k)].width = 13
pr.freeze_panes = 'G6'

ch = LineChart()
ch.title = 'Valor neto total, año a año'
ch.height, ch.width = 9, 26
ch.y_axis.numFmt = '#,##0'
ch.add_data(Reference(pr, min_col=4, min_row=5, max_row=RN), titles_from_data=True)
ch.set_categories(Reference(pr, min_col=1, min_row=R0, max_row=RN))
pr.add_chart(ch, openpyxl.utils.get_column_letter(7 + NCAR * 6 + 1) + '6')


# =========================================================================
#  HOJA 0 - PANEL  (lo único que se toca para jugar)
# =========================================================================
pa = wb.create_sheet('Panel', wb.sheetnames.index('Objetivos'))
titulo(pa, 'A1', 'PANEL DE MANDO')
nota(pa, 'A2', 'Toca SOLO las celdas verdes de esta pestaña. Todo lo demás del libro se recalcula. '
               'Las pestañas «Objetivos» y «Proyección» son el detalle: no hace falta abrirlas para jugar.')

pa['A4'] = '=Proyección!A3'
pa['A4'].font = Font(bold=True, size=16, color='FFC00000')
pa.row_dimensions[4].height = 26
pa['A5'] = '=IF(LEFT(Proyección!A3,1)="✓","Con estos números tu plan se sostiene.",\
"Sube la aportación de la cartera que salga en rojo abajo, o retrasa/abarata ese objetivo.")'
pa['A5'].font = Font(size=10, italic=True, color='FF808080')

banda(pa, 7, list('ABCDEF'), 'TUS CARTERAS  — lo que aportas a cada una')
cabecera(pa, 8, [('A', 'CARTERA'), ('B', 'RENTABILIDAD'), ('C', 'APORTAS\n€/mes'),
                 ('D', 'NECESITA\n€/mes'), ('E', 'DIFERENCIA'), ('F', 'ESTADO')], alto=30)
for i, (nom, rent, inv, extra, mens, nt) in enumerate(CARTERAS):
    r = 9 + i
    cr = CAR0 + i
    pa[f'A{r}'] = nom
    pa[f'A{r}'].font = Font(bold=True, size=11, color=AZUL_OSC)
    pa[f'B{r}'] = rent
    pa[f'B{r}'].number_format = PCT
    editable(pa[f'B{r}'])
    if mens is None:
        pa[f'C{r}'] = f'={CELDA_PP_ANUAL}/12'
        pa[f'C{r}'].font = Font(bold=True, size=11, color=AZUL_OSC)
    else:
        pa[f'C{r}'] = mens
        editable(pa[f'C{r}'])
    pa[f'C{r}'].number_format = EUR0
    pa[f'D{r}'] = f'=Objetivos!F{cr}'
    pa[f'E{r}'] = f'=C{r}-D{r}'
    for col in 'DE':
        pa[f'{col}{r}'].number_format = EUR0
        pa[f'{col}{r}'].font = Font(bold=True, size=11, color=AZUL_OSC)
    pa[f'F{r}'] = f'=Objetivos!G{cr}'
    pa[f'F{r}'].font = Font(bold=True, size=11, color='FFC00000')
    pa.row_dimensions[r].height = 20
rr = 9 + NCAR
pa[f'A{rr}'] = 'TOTAL'
for col in 'ABCDEF':
    pa[f'{col}{rr}'].fill = PatternFill('solid', fgColor=NARANJA)
    pa[f'{col}{rr}'].font = Font(bold=True, size=12, color=AZUL_OSC)
for col in 'CDE':
    pa[f'{col}{rr}'] = f'=SUM({col}9:{col}{rr - 1})'
    pa[f'{col}{rr}'].number_format = EUR0

banda(pa, 16, list('ABC'), 'SUPUESTOS')
SUP = [(17, 'Inflación anual', INF, PCT, 'Cuánto sube el precio de las cosas'),
       (18, 'Impuestos sobre la plusvalía', IMP, PCT, 'Tipo medio al vender (en España la base del ahorro va por tramos, del 19% al 30%)'),
       (19, 'Incremento anual de la aportación', INCR, PCT, 'Si es menor que la inflación, cada año inviertes menos en términos reales'),
       (20, 'Dejo de aportar dentro de (años)', None, '0', 'Vacío = sigue al objetivo de tipo «Jubilación»')]
for row, etiq, val, fmt, ayuda in SUP:
    pa[f'A{row}'] = etiq
    pa[f'A{row}'].font = Font(bold=True, size=10, color=AZUL_OSC)
    if val is not None:
        pa[f'B{row}'] = val
    pa[f'B{row}'].number_format = fmt
    editable(pa[f'B{row}'])
    pa[f'C{row}'] = ayuda
    pa[f'C{row}'].font = Font(size=9, color='FF808080')

banda(pa, 22, list('ABC'), 'CÓMO VAS')
RESU = [(23, 'Aportas al mes (todas las carteras)', '=C' + str(rr), EUR, ''),
        (24, 'Necesitas al mes', '=D' + str(rr), EUR, 'Suma de lo que pide cada objetivo'),
        (25, '→ TE FALTA / TE SOBRA', f'=C{rr}-D{rr}', EUR, 'En negativo, es lo que hay que subir'),
        (26, 'Tu ahorro real (de tus cuentas)', '=Objetivos!B13', EUR, 'Lo que de verdad te sobra cada mes'),
        (27, '% de tu ahorro que estás invirtiendo', '=Objetivos!B15', PCT, 'Sin el plan de pensiones'),
        (28, 'Dejas de aportar en', '=Objetivos!B12', '0', 'El año de tu jubilación'),
        (29, 'Patrimonio neto al jubilarte',
         "=IFERROR(INDEX(Proyección!D:D,MATCH(Objetivos!B12,Proyección!A:A,0)),0)", EUR,
         'Lo que tendrás, ya descontados los impuestos latentes')]
for row, etiq, form, fmt, ayuda in RESU:
    pa[f'A{row}'] = etiq
    pa[f'A{row}'].font = Font(bold=True, size=10, color=AZUL_OSC)
    pa[f'B{row}'] = form
    pa[f'B{row}'].number_format = fmt
    pa[f'B{row}'].font = Font(bold=True, size=12, color=AZUL_OSC)
    pa[f'C{row}'] = ayuda
    pa[f'C{row}'].font = Font(size=9, color='FF808080')
for col in 'ABC':
    pa[f'{col}25'].fill = PatternFill('solid', fgColor=VERDE)
pa['B25'].font = Font(bold=True, size=14, color='FF006100')

for col, w in (('A', 38), ('B', 16), ('C', 16), ('D', 14), ('E', 14), ('F', 24)):
    pa.column_dimensions[col].width = w

ch2 = LineChart()
ch2.title = 'Valor neto total, año a año'
ch2.height, ch2.width = 9, 24
ch2.y_axis.numFmt = '#,##0'
ch2.add_data(Reference(pr, min_col=4, min_row=5, max_row=RN), titles_from_data=True)
ch2.set_categories(Reference(pr, min_col=1, min_row=R0, max_row=RN))
pa.add_chart(ch2, 'A32')

# --- chequeo: nada que la Proyeccion sume puede devolver texto
malas = [f'{c}{r}: {ob[f"{c}{r}"].value}' for r in range(OBJ0, OBJN + 1) for c in ('E', 'H')
         if isinstance(ob[f'{c}{r}'].value, str) and ob[f'{c}{r}'].value.startswith('=')
         and '""' in ob[f'{c}{r}'].value.split(',', 1)[-1]]
if malas:
    raise SystemExit('ABORTADO — pueden devolver "" y la Proyección las suma:\n  ' + '\n  '.join(malas))
print('chequeo aritmética Objetivos↔Proyección: OK')

# --- chequeo: cada cartera de Objetivos apunta a SU fila del Panel
for i, (nom, *_r) in enumerate(CARTERAS):
    fila_panel = 9 + i
    if pa[f'A{fila_panel}'].value != nom:
        raise SystemExit(f'ABORTADO — Objetivos!B{CAR0+i} apunta a Panel fila {fila_panel} '
                         f'({pa[f"A{fila_panel}"].value!r}) pero la cartera es {nom!r}')
    for col_ob, col_pa in (('B', 'B'), ('E', 'C')):
        esperado = f'=Panel!{col_pa}{fila_panel}'
        real = ob[f'{col_ob}{CAR0 + i}'].value
        if real != esperado:
            raise SystemExit(f'ABORTADO — Objetivos!{col_ob}{CAR0+i} = {real!r}, esperaba {esperado!r}')
print('chequeo referencias Panel↔Objetivos: OK')

wb.save(OUT)
print('libro escrito ->', OUT)

# =========================================================================
#  El MISMO modelo en Python
# =========================================================================
wv = openpyxl.load_workbook(CALC, data_only=True)
ANIO0 = wv['Patrimonio Neto']['E1'].value.year
def leer(ref):
    hoja, celda = ref.rsplit('!', 1)
    return wv[hoja.strip("'")][celda].value
PP_MES = leer(CELDA_PP_ANUAL) / 12
AHORRO = leer(CELDA_AHORRO)

CAR = []
for i, (nom, rent, inv, extra, mens) in enumerate([(c[0], c[1], c[2], c[3], c[4]) for c in CARTERAS]):
    yaInv = leer(inv) if inv else 0.0
    CAR.append({'nom': nom, 'rent': rent, 'ini': yaInv + extra,
                'inv': yaInv, 'extra': extra,
                'mens': PP_MES if mens is None else mens, 'row': CAR0 + i})
AFIN = min([o[3] - ANIO0 for o in OBJ if o[2] == 'Jubilación'] or [999])
V = {}


def p(ref, val):
    V[('Objetivos', ref)] = val


p('B5', ANIO0)
p('B8', INF / (1 - IMP))
p('B11', AFIN)
p('B12', ANIO0 + AFIN)
p('B13', AHORRO)
p('B14', sum(c['mens'] for c in CAR))
p('B15', (sum(c['mens'] for c in CAR) - CAR[2]['mens']) / AHORRO if AHORRO else 0)
for c in CAR:
    p(f'C{c["row"]}', c['inv'])
    p(f'E{c["row"]}', c['mens'])
for col, vals in (('C', [c['inv'] for c in CAR]), ('D', [c['extra'] for c in CAR]),
                  ('E', [c['mens'] for c in CAR])):
    p(f'{col}{CARN + 1}', sum(vals))

for r in range(OBJ0, OBJN + 1):
    p(f'N{r}', 9999)
    p(f'E{r}', 0)
por_cartera = {c['nom']: {'unicos': {}, 'rentas': []} for c in CAR}
nec = {c['nom']: 0.0 for c in CAR}
tot = {'I': 0.0, 'J': 0.0, 'K': 0.0, 'L': 0.0}
rentab = {c['nom']: c['rent'] for c in CAR}
for i, (nom, car, tipo, anio_abs, imp, renta, cons, nt) in enumerate(OBJ):
    r = OBJ0 + i
    if tipo is None:
        for col in 'IJKL':
            p(f'{col}{r}', '')
        continue
    d = anio_abs - ANIO0
    p(f'E{r}', d)
    ro = rentab.get(car, INF / (1 - IMP))
    renta = renta or 0          # celda vacía = 0, como en el Excel
    imp = imp or 0
    cons = cons or 0
    if tipo == 'Jubilación':
        p(f'N{r}', d)
    if tipo in ('Renta', 'Jubilación'):
        I_ = renta * 12 * (1 + INF) ** d * ((1 + INF) ** cons - 1) / INF
        q = (1 + INF) / (1 + ro)
        fac = cons if abs(INF - ro) < 1e-7 else (1 - q ** cons) / (1 - q)
        J_ = renta * 12 * (1 + INF) ** d * fac
        por_cartera[car]['rentas'].append((d, cons, renta))
    else:
        I_ = imp * (1 + INF) ** d
        J_ = I_
        por_cartera[car]['unicos'][d] = por_cartera[car]['unicos'].get(d, 0.0) + I_
    K_ = J_ / (1 + ro) ** d
    L_ = ((J_ * ro) / ((1 + ro) ** d - 1)) / 12 if d > 0 else ''
    for col, val in (('I', I_), ('J', J_), ('K', K_), ('L', L_)):
        p(f'{col}{r}', val)
        if isinstance(val, float):
            tot[col] += val
    if isinstance(L_, float):
        nec[car] += L_
for col in 'IJKL':
    p(f'{col}{TOT}', tot[col])
for c in CAR:
    p(f'F{c["row"]}', nec[c['nom']])
    est = ('— sin objetivos' if nec[c['nom']] == 0 else
           ('✓ cubierta' if c['mens'] >= nec[c['nom']] else f'⚠ faltan {nec[c["nom"]] - c["mens"]:,.0f} €/mes'
            .replace(',', '.')))
    p(f'G{c["row"]}', est)
p(f'F{CARN + 1}', sum(nec.values()))

bolsas = [{'val': 0.0, 'base': 0.0, 'coste': 0.0} for _ in CAR]
filas = []
for i in range(ANIOS):
    a = i + 1
    r = R0 + i
    V[('Proyección', f'A{r}')] = ANIO0 + a
    tot_val = tot_neto = tot_ap = tot_re = tot_sc = 0.0
    for k, c in enumerate(CAR):
        b = bolsas[k]
        if a >= AFIN:
            ap = 0.0
        elif a == 1:
            ap = c['mens'] * 12 + c['ini']
        else:
            ap = c['mens'] * 12 * (1 + INCR) ** (a - 1)
        if a == 1:
            val = ap
            base = ap
        else:
            val = max(0.0, (b['val'] - b['coste']) * (1 + c['rent']) + ap)
            base = max(0.0, b['base'] - (b['re'] * b['base'] / b['val'] if b['val'] else 0.0) + ap)
        re = por_cartera[c['nom']]['unicos'].get(a, 0.0)
        for (ini_a, cons, renta) in por_cartera[c['nom']]['rentas']:
            if ini_a <= a < ini_a + cons:
                re += renta * 12 * (1 + INF) ** a
        pedido_py = re
        re = min(val, re)
        sin_cubrir = max(0.0, pedido_py - val)
        coste = re + (re * max(0.0, val - base) / val if val else 0.0) * IMP
        bolsas[k] = {'val': val, 'base': base, 're': re, 'coste': coste, 'sc': sin_cubrir}
        tot_sc = tot_sc + sin_cubrir
        cb = 7 + k * 6
        for j, v in enumerate((ap, val, base, re, coste, sin_cubrir)):
            V[('Proyección', f'{openpyxl.utils.get_column_letter(cb + j)}{r}')] = v
        tot_val += val
        tot_neto += val - max(0.0, val - base) * IMP
        tot_ap += ap
        tot_re += re
    for col, v in (('C', tot_val), ('D', tot_neto), ('E', tot_ap), ('F', tot_re)):
        V[('Proyección', f'{col}{r}')] = v
    filas.append((a, ANIO0 + a, tot_ap, tot_val, tot_neto, tot_re,
                  [bolsas[k]['val'] for k in range(NCAR)], tot_sc))

agot = [y for (a, y, ap, v, n, re, bb, sc) in filas if sc > 0.01]
V[('Proyección', 'A3')] = (f'⚠ NO LLEGAS: el primer objetivo que se queda sin pagar es en {agot[0]}'
                           if agot else '✓ Todos los objetivos se pagan')
jub_row = R0 + AFIN - 1
V[('Panel', 'A4')] = V[('Proyección', 'A3')]
V[('Panel', 'A5')] = ('Con estos números tu plan se sostiene.' if V[('Proyección', 'A3')].startswith('✓')
                      else 'Sube la aportación de la cartera que salga en rojo abajo, '
                           'o retrasa/abarata ese objetivo.')
for i, c in enumerate(CAR):
    V[('Panel', f'D{9+i}')] = nec[c['nom']]
    V[('Panel', f'E{9+i}')] = c['mens'] - nec[c['nom']]
    V[('Panel', f'F{9+i}')] = V[('Objetivos', 'G%d' % c['row'])]
    if c['mens'] == PP_MES and CARTERAS[i][4] is None:
        V[('Panel', f'C{9+i}')] = PP_MES
V[('Panel', f'C{9+NCAR}')] = sum(c['mens'] for c in CAR)
V[('Panel', f'D{9+NCAR}')] = sum(nec.values())
V[('Panel', f'E{9+NCAR}')] = sum(c['mens'] for c in CAR) - sum(nec.values())
V[('Panel', 'B23')] = sum(c['mens'] for c in CAR)
V[('Panel', 'B24')] = sum(nec.values())
V[('Panel', 'B25')] = sum(c['mens'] for c in CAR) - sum(nec.values())
V[('Panel', 'B26')] = AHORRO
V[('Panel', 'B27')] = V[('Objetivos', 'B15')]
V[('Panel', 'B28')] = ANIO0 + AFIN
V[('Panel', 'B29')] = next((n for (a, y, ap, v, n, re, bb, sc) in filas if y == ANIO0 + AFIN), 0)
json.dump({f'{s}||{c}': v for (s, c), v in V.items()}, open(VALJSON, 'w'), ensure_ascii=False)

print()
print('=== CARTERAS ===')
print(f'  {"cartera":20s} {"rent.":>7} {"ya inv.":>10} {"extra":>9} {"mensual":>9} {"necesario":>11}   estado')
for c in CAR:
    est_c = V[('Objetivos', 'G%d' % c['row'])]
    print(f'  {c["nom"]:20s} {c["rent"]*100:>6.2f}% {c["inv"]:>10,.0f} {c["extra"]:>9,.0f} '
          f'{c["mens"]:>9,.0f} {nec[c["nom"]]:>11,.0f}   {est_c}')
print()
print('=== OBJETIVOS ===')
print(f'  {"objetivo":20s} {"cartera":16s} {"año":>5} {"cap. nec.":>12} {"mens. nec.":>11}')
for i, (nom, car, tipo, anio_abs, imp, renta, cons, nt) in enumerate(OBJ):
    r = OBJ0 + i
    if V[('Objetivos', f'I{r}')] == '':
        print(f'  {nom:20s} {"":16s} {anio_abs:>5}   — falta info')
        continue
    print(f'  {nom:20s} {car:16s} {anio_abs:>5} {V[("Objetivos", f"J{r}")]:>12,.0f} '
          f'{V[("Objetivos", f"L{r}")]:>11,.0f}')
print(f'  {"TOTAL":20s} {"":16s} {"":>5} {tot["J"]:>12,.0f} {tot["L"]:>11,.0f}')
print()
print('=== PROYECCIÓN ===')
print(f'  {"año":>5} {"aportas":>9} ' + ' '.join(f'{c["nom"][:12]:>13}' for c in CAR) +
      f' {"TOTAL":>12} {"retiras":>10}')
for (a, y, ap, v, n, re, bb, sc) in filas:
    if re > 0 or a in (1, 2) or a % 10 == 0 or a == AFIN - 1:
        print(f'  {y:>5} {ap:>9,.0f} ' + ' '.join(f'{x:>13,.0f}' for x in bb) + f' {v:>12,.0f} {re:>10,.0f}')
    if a > 62:
        break
print()
print('DIAGNÓSTICO:', V[('Proyección', 'A3')])
