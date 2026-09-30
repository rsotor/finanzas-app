# -*- coding: utf-8 -*-
"""Reconstruye Finanzas personales.xlsx con la estructura acordada."""
import os, json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import PieChart, Reference

EUR = '#,##0.00\\ "€"'
EUR0 = '#,##0.00\\ "€";-#,##0.00\\ "€";""'
PCT0 = '0.0%;-0.0%;""'
PCT = '0.0%'
AZUL   = "1F3864"; AZUL_CL = "D9E2F3"; GRIS = "F2F2F2"; VERDE = "E2EFDA"; NARANJA = "FCE4D6"
TH = Side(style="thin", color="BFBFBF")
BORDE = Border(bottom=TH)

wb = openpyxl.Workbook()
ws = wb.active; ws.title = "Ingresos - Gastos"

# ---------------------------------------------------------------- datos
ING = [
 ("Laborales", [
   ("Nómina neta Ana", 2100.00, None,
    "EJEMPLO. Lo que entra en la cuenta cada mes (neto). Si cobras pagas extra prorrateadas, pon el neto mensual; si no, pon el total anual en la columna C"),
   ("Nómina neta Luis", 1850.00, None,
    "EJEMPLO. Mismo criterio que la fila de arriba"),
   ("Retribución flexible (seguro médico de empresa)", 90.00, None,
    "EJEMPLO. Si la empresa te paga algo en especie (seguro, ticket, plan de pensiones), entra como ingreso y sale como gasto: así el total es honesto"),
   ("Ingresos extra (variable)", None, 1200.00,
    "EJEMPLO. Freelance, bonus, ventas de segunda mano… Anual y sin inflar: solo lo que de verdad cobras"),
   (None, None, None, None),
 ]),
 ("Inversiones", [
   ("Intereses cuentas y depósitos", None, 300.00,
    "EJEMPLO. Intereses brutos al año de la cuenta remunerada (antes de la retención de Hacienda)"),
   (None, None, None, None),
   (None, None, None, None),
 ]),
]

GAS = [
 ("INDISPENSABLES", [
   ("Vivienda habitual", [
     ("Hipoteca / alquiler", 750.00, None, "EJEMPLO. Cuota mensual"),
     ("Comunidad", 70.00, None, "EJEMPLO"),
     ("Seguro vivienda", None, 320.00, "EJEMPLO. Recibo anual"),
     ("IBI", None, 450.00, "EJEMPLO. Recibo anual del ayuntamiento"),
     ("Tasa basuras", None, 110.00, "EJEMPLO"),
     (None, None, None, None), (None, None, None, None),
   ]),
   ("Seguros", [
     ("Seguro de vida", None, 360.00, "EJEMPLO. Suele ir ligado a la hipoteca: revisa si cubre solo la deuda o también a la familia"),
     ("Seguro de salud", 90.00, None, "EJEMPLO. Si te lo paga la empresa, aparece también arriba como retribución flexible"),
     (None, None, None, None), (None, None, None, None),
   ]),
   ("Ahorro / Inversión", [
     ("Jubilación (plan de pensiones)", 150.00, None, "EJEMPLO. Aportación mensual al plan. No es dinero disponible: por eso va aquí como gasto y no en el ahorro del final. La pestaña «Panel» la lee de esta fila"),
     (None, None, None, None), (None, None, None, None),
   ]),
 ]),
 ("NECESARIOS", [
   ("Recibos y servicios del hogar", [
     ("Agua", None, 240.00, "EJEMPLO. Suma de los recibos del año"),
     ("Luz", 55.00, None, "EJEMPLO"),
     ("Gas / calefacción", None, 420.00, "EJEMPLO. Muy estacional: mejor el total anual"),
     ("Teléfono / Internet", 45.00, None, "EJEMPLO"),
     ("Mantenimiento del hogar", None, 250.00, "EJEMPLO. Revisión de la caldera, pequeñas reparaciones"),
     (None, None, None, None), (None, None, None, None),
   ]),
   ("Alimentación", [
     ("Supermercado", 520.00, None, "EJEMPLO. Truco: suma los cargos de supermercados de 3 meses del extracto y divide entre 3"),
     ("Fruterías / mercado", 60.00, None, "EJEMPLO"),
     (None, None, None, None), (None, None, None, None),
   ]),
   ("Cuidado personal / familiar", [
     ("Ropa y calzado", None, 1100.00, "EJEMPLO"),
     ("Farmacia / médicos", None, 250.00, "EJEMPLO"),
     ("Peluquería y cosmética", None, 300.00, "EJEMPLO"),
     (None, None, None, None), (None, None, None, None),
   ]),
   ("Educación / cuidado de los hijos", [
     ("Guardería / colegio", 280.00, None, "EJEMPLO"),
     ("Comedor escolar", None, 1300.00, "EJEMPLO. Curso completo (normalmente no se cobra en verano)"),
     ("Libros, material y uniforme", None, 250.00, "EJEMPLO"),
     ("Actividades extraescolares", 40.00, None, "EJEMPLO"),
     (None, None, None, None), (None, None, None, None),
   ]),
   ("Transporte", [
     ("Gasolina", 110.00, None, "EJEMPLO"),
     ("Seguro del coche", None, 450.00, "EJEMPLO. Recibo anual"),
     ("Mantenimiento coche (ITV, taller, ruedas)", None, 400.00, "EJEMPLO"),
     ("Impuesto de circulación", None, 70.00, "EJEMPLO"),
     ("Transporte público / parking", 20.00, None, "EJEMPLO"),
     (None, None, None, None), (None, None, None, None),
   ]),
 ]),
 ("REVISABLES", [
   ("Comer y beber fuera", [
     ("Restaurantes / bares / cafés", 120.00, None, "EJEMPLO. Es de lo primero que se recorta en una emergencia"),
     (None, None, None, None),
   ]),
   ("Suscripciones", [
     ("Plataformas de vídeo", 18.00, None, "EJEMPLO"),
     ("Música", 11.00, None, "EJEMPLO"),
     ("Almacenamiento en la nube / apps", None, 30.00, "EJEMPLO"),
     ("Compras online con cuota anual", None, 50.00, "EJEMPLO"),
     (None, None, None, None), (None, None, None, None),
   ]),
   ("Ocio y tiempo libre", [
     ("Ocio (cine, parques, excursiones)", 60.00, None, "EJEMPLO"),
     ("Vacaciones", None, 1800.00, "EJEMPLO"),
     ("Regalos (cumpleaños, Navidad)", None, 500.00, "EJEMPLO"),
     (None, None, None, None), (None, None, None, None),
   ]),
   ("Otros revisables", [
     ("Gastos extra e imprevistos", None, 900.00, "EJEMPLO. Partida para lo puntual: un electrodoméstico, una avería. Mejor presupuestarlo que llevarse la sorpresa"),
     ("Compras varias online", 40.00, None, "EJEMPLO"),
     (None, None, None, None), (None, None, None, None),
   ]),
 ]),
]

# ---------------------------------------------------------------- helpers
def titulo(ws, r, txt, sub=None):
    c = ws.cell(r, 1, txt); c.font = Font(bold=True, size=15, color=AZUL)
    if sub:
        c2 = ws.cell(r+1, 1, sub); c2.font = Font(size=9, italic=True, color="808080")

def cabecera(ws, r, txt, fill):
    heads = [txt, "MENSUAL", "ANUAL", "TOTAL ANUAL", "%", "", "Origen del dato · avisos"]
    for j, h in enumerate(heads, start=1):
        c = ws.cell(r, j, h)
        c.font = Font(bold=True, size=11, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="left" if j in (1, 7) else "center", vertical="center")
    ws.row_dimensions[r].height = 20

def linea_grupo(ws, r, nombre, first, last, tot_ref, nivel=1):
    c = ws.cell(r, 1, ("   " if nivel == 2 else "") + nombre)
    c.font = Font(bold=True, size=11 if nivel == 1 else 10, color=AZUL)
    fill = AZUL_CL if nivel == 1 else GRIS
    for j in range(1, 6):
        ws.cell(r, j).fill = PatternFill("solid", fgColor=fill)
        ws.cell(r, j).border = BORDE
    ws.cell(r, 2, f"=SUM(B{first}:B{last})").number_format = EUR
    ws.cell(r, 3, f"=SUM(C{first}:C{last})").number_format = EUR
    ws.cell(r, 4, f"=SUM(D{first}:D{last})").number_format = EUR
    ws.cell(r, 5, f"=IFERROR(D{r}/{tot_ref},0)").number_format = PCT
    for j in (2, 3, 4, 5):
        ws.cell(r, j).font = Font(bold=True, size=10, color=AZUL)

def linea_concepto(ws, r, nombre, mens, anual, nota, tot_ref):
    ws.cell(r, 1, nombre if nombre else None).font = Font(size=10)
    ws.cell(r, 1).alignment = Alignment(indent=2)
    if mens is not None: ws.cell(r, 2, mens)
    if anual is not None: ws.cell(r, 3, anual)
    ws.cell(r, 2).number_format = EUR0
    ws.cell(r, 3).number_format = EUR0
    ws.cell(r, 4, f"=IF(B{r}>0,B{r}*12,C{r})").number_format = EUR0
    ws.cell(r, 5, f"=IFERROR(D{r}/{tot_ref},0)").number_format = PCT0
    if nota:
        n = ws.cell(r, 7, nota); n.font = Font(size=8, color="808080")
        n.alignment = Alignment(vertical="top", wrap_text=False)
    for j in range(1, 6):
        ws.cell(r, j).border = BORDE

def linea_total(ws, r, txt, formula, fill=VERDE, pct=None):
    c = ws.cell(r, 1, txt); c.font = Font(bold=True, size=12, color=AZUL)
    ws.cell(r, 4, formula).number_format = EUR
    ws.cell(r, 4).font = Font(bold=True, size=12, color=AZUL)
    if pct: ws.cell(r, 5, pct).number_format = PCT
    for j in range(1, 6):
        ws.cell(r, j).fill = PatternFill("solid", fgColor=fill)
        ws.cell(r, j).border = Border(top=Side(style="medium", color=AZUL), bottom=Side(style="medium", color=AZUL))

# ---------------------------------------------------------------- INGRESOS - GASTOS
titulo(ws, 1, "INGRESOS Y GASTOS — Familia de ejemplo (Ana, Luis y Leo)",
       "CIFRAS INVENTADAS. Sustitúyelas por las tuyas: columna B si el gasto es mensual, columna C si es anual (nunca las dos). La columna G es para apuntar de dónde sale cada dato.")
R_TOT_ING = None
FILA = {}          # concepto -> fila en «Ingresos - Gastos»
mapa = []          # (tipo, fila, etiqueta) para construir Presupuesto

r = 4
cabecera(ws, r, "INGRESOS", AZUL); FILA_CAB_ING = r; r += 1
grupos_ing = []
for nombre, items in ING:
    fila_g = r; r += 1
    first = r
    for it in items:
        linea_concepto(ws, r, it[0], it[1], it[2], it[3], "$D$%d")
        if it[0]: FILA[it[0]] = r
        r += 1
    last = r - 1
    grupos_ing.append((fila_g, nombre, first, last))
r_tot_ing = r
for fila_g, nombre, first, last in grupos_ing:
    linea_grupo(ws, fila_g, nombre, first, last, f"$D${r_tot_ing}")
    for rr in range(first, last + 1):
        ws.cell(rr, 5, f"=IFERROR(D{rr}/$D${r_tot_ing},0)").number_format = PCT
linea_total(ws, r_tot_ing, "TOTAL INGRESOS", "=" + "+".join(f"D{g[0]}" for g in grupos_ing))
r += 2

# GASTOS
cabecera(ws, r, "GASTOS", AZUL); r += 1
bloques = []
for bloque, subgrupos in GAS:
    fila_b = r; r += 1
    subs = []
    for sub, items in subgrupos:
        fila_s = r; r += 1
        first = r
        for it in items:
            linea_concepto(ws, r, it[0], it[1], it[2], it[3], "X")
            if it[0]: FILA[it[0]] = r
            r += 1
        subs.append((fila_s, sub, first, r - 1))
    bloques.append((fila_b, bloque, subs))
    r += 1
r_tot_gas = r
TG = f"$D${r_tot_gas}"
for fila_b, bloque, subs in bloques:
    ws.cell(fila_b, 1, bloque).font = Font(bold=True, size=12, color="FFFFFF")
    for j in range(1, 6):
        ws.cell(fila_b, j).fill = PatternFill("solid", fgColor="4472C4")
    ws.cell(fila_b, 4, "=" + "+".join(f"D{s[0]}" for s in subs)).number_format = EUR
    ws.cell(fila_b, 4).font = Font(bold=True, size=12, color="FFFFFF")
    ws.cell(fila_b, 5, f"=IFERROR(D{fila_b}/{TG},0)").number_format = PCT
    ws.cell(fila_b, 5).font = Font(bold=True, size=11, color="FFFFFF")
    ws.row_dimensions[fila_b].height = 18
    for fila_s, sub, first, last in subs:
        linea_grupo(ws, fila_s, sub, first, last, TG, nivel=2)
        for rr in range(first, last + 1):
            ws.cell(rr, 5, f"=IFERROR(D{rr}/{TG},0)").number_format = PCT0
linea_total(ws, r_tot_gas, "TOTAL GASTOS", "=" + "+".join(f"D{b[0]}" for b in bloques), fill=NARANJA)

# --- semaforo: cuanto mas gordo el gasto, mas rojo
from openpyxl.formatting.rule import ColorScaleRule
rangos = []
for _fb, _bl, _subs in bloques:
    for _fs, _sub, _fi, _la in _subs:
        rangos.append(f"E{_fi}:E{_la}")
        rangos.append(f"D{_fi}:D{_la}")
esc_e = " ".join(r for r in rangos if r.startswith("E"))
esc_d = " ".join(r for r in rangos if r.startswith("D"))
for rng in (esc_e, esc_d):
    ws.conditional_formatting.add(rng, ColorScaleRule(
        start_type="num",   start_value=0,   start_color="63BE7B",
        mid_type="percentile", mid_value=70,  mid_color="FFEB84",
        end_type="max",                        end_color="F8696B"))
# los subtotales de bloque tambien, con su propia escala
sub_e = " ".join(f"E{_fs}" for _fb, _bl, _subs in bloques for _fs, _sub, _fi, _la in _subs)
ws.conditional_formatting.add(sub_e, ColorScaleRule(
    start_type="min", start_color="C6EFCE", mid_type="percentile", mid_value=50, mid_color="FFEB84",
    end_type="max", end_color="FFC7CE"))
r = r_tot_gas + 2
ws.cell(r, 1, "AHORRO ANUAL").font = Font(bold=True, size=13, color="006100")
ws.cell(r, 4, f"=D{r_tot_ing}-D{r_tot_gas}").number_format = EUR
ws.cell(r, 4).font = Font(bold=True, size=13, color="006100")
ws.cell(r, 7, "Ojo: la aportación al plan de pensiones (bloque Ahorro/Inversión) ya está contada como gasto, así que NO está aquí dentro.").font = Font(size=8, color="808080")
ws.cell(r+1, 1, "AHORRO MENSUAL").font = Font(bold=True, size=13, color="006100")
ws.cell(r+1, 4, f"=D{r}/12").number_format = EUR
ws.cell(r+1, 4).font = Font(bold=True, size=13, color="006100")
ws.cell(r+2, 1, "Tasa de ahorro").font = Font(bold=True, size=11, color="006100")
ws.cell(r+2, 4, f"=IFERROR(D{r}/D{r_tot_ing},0)").number_format = PCT
ws.cell(r+2, 4).font = Font(bold=True, size=11, color="006100")
for rr in (r, r+1, r+2):
    for j in range(1, 6):
        ws.cell(rr, j).fill = PatternFill("solid", fgColor=VERDE)
R_AHORRO = r

ws.column_dimensions["A"].width = 52
for col in ("B", "C", "D", "E"): ws.column_dimensions[col].width = 14
ws.column_dimensions["F"].width = 3
ws.column_dimensions["G"].width = 110
ws.freeze_panes = "A5"
ws.sheet_view.showGridLines = False

# ---------------------------------------------------------------- PRESUPUESTO
wp = wb.create_sheet("Presupuesto")
IG = "'Ingresos - Gastos'"
titulo(wp, 1, "PRESUPUESTO",
       "Izquierda: lo que gastas hoy (viene solo de «Ingresos - Gastos», no se toca). Derecha: lo que QUIERES gastar — edita la columna D. La diferencia te dice cuánto liberas.")
r = 4
heads = ["CONCEPTO", "ACTUAL MENSUAL", "ACTUAL ANUAL", "PRESUPUESTO MENSUAL", "PRESUPUESTO ANUAL", "DIFERENCIA ANUAL", "%"]
for j, h in enumerate(heads, start=1):
    c = wp.cell(r, j, h)
    c.font = Font(bold=True, size=10, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=AZUL)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
wp.row_dimensions[r].height = 30
r += 1

p_bloques = []
for fila_b, bloque, subs in bloques:
    fb = r
    wp.cell(r, 1, bloque).font = Font(bold=True, size=12, color="FFFFFF")
    for j in range(1, 8): wp.cell(r, j).fill = PatternFill("solid", fgColor="4472C4")
    r += 1
    p_subs = []
    for fila_s, sub, first, last in subs:
        fs = r
        wp.cell(r, 1, "   " + sub).font = Font(bold=True, size=10, color=AZUL)
        for j in range(1, 8): wp.cell(r, j).fill = PatternFill("solid", fgColor=GRIS)
        r += 1
        pf = r
        for rr in range(first, last + 1):
            wp.cell(r, 1, f'=IF({IG}!A{rr}="","",{IG}!A{rr})').alignment = Alignment(indent=2)
            wp.cell(r, 1).font = Font(size=10)
            wp.cell(r, 2, f"={IG}!D{rr}/12").number_format = EUR0
            wp.cell(r, 3, f"={IG}!D{rr}").number_format = EUR0
            wp.cell(r, 4, f"=B{r}").number_format = EUR0          # editable
            wp.cell(r, 4).font = Font(size=10, color="0070C0", bold=True)
            wp.cell(r, 5, f"=D{r}*12").number_format = EUR0
            wp.cell(r, 6, f"=C{r}-E{r}").number_format = EUR0
            for j in range(1, 8): wp.cell(r, j).border = BORDE
            r += 1
        pl = r - 1
        for col in (2, 3, 4, 5, 6):
            wp.cell(fs, col, f"=SUM({get_column_letter(col)}{pf}:{get_column_letter(col)}{pl})").number_format = EUR
            wp.cell(fs, col).font = Font(bold=True, size=10, color=AZUL)
        p_subs.append(fs)
    for col in (2, 3, 4, 5, 6):
        wp.cell(fb, col, "=" + "+".join(f"{get_column_letter(col)}{x}" for x in p_subs)).number_format = EUR
        wp.cell(fb, col).font = Font(bold=True, size=12, color="FFFFFF")
    p_bloques.append(fb)
    r += 1

R_PRES_TOT = r
wp.cell(r, 1, "TOTAL GASTOS").font = Font(bold=True, size=12, color=AZUL)
for col in (2, 3, 4, 5, 6):
    wp.cell(r, col, "=" + "+".join(f"{get_column_letter(col)}{x}" for x in p_bloques)).number_format = EUR
    wp.cell(r, col).font = Font(bold=True, size=12, color=AZUL)
for j in range(1, 8): wp.cell(r, j).fill = PatternFill("solid", fgColor=NARANJA)
for fb in p_bloques:
    wp.cell(fb, 7, f"=IFERROR(E{fb}/$E${R_PRES_TOT},0)").number_format = PCT
    wp.cell(fb, 7).font = Font(bold=True, size=11, color="FFFFFF")
r += 2
wp.cell(r, 1, "AHORRO ANUAL PRESUPUESTADO").font = Font(bold=True, size=12, color="006100")
wp.cell(r, 5, f"={IG}!D{r_tot_ing}-E{R_PRES_TOT}").number_format = EUR
wp.cell(r, 5).font = Font(bold=True, size=12, color="006100")
wp.cell(r+1, 1, "AHORRO MENSUAL PRESUPUESTADO").font = Font(bold=True, size=12, color="006100")
wp.cell(r+1, 5, f"=E{r}/12").number_format = EUR
wp.cell(r+1, 5).font = Font(bold=True, size=12, color="006100")
wp.cell(r+2, 1, "MEJORA FRENTE A HOY (anual)").font = Font(bold=True, size=11, color="006100")
wp.cell(r+2, 5, f"=F{R_PRES_TOT}").number_format = EUR
wp.cell(r+2, 5).font = Font(bold=True, size=11, color="006100")
for rr in (r, r+1, r+2):
    for j in range(1, 8): wp.cell(rr, j).fill = PatternFill("solid", fgColor=VERDE)
wp.column_dimensions["A"].width = 52
for col in ("B", "C", "D", "E", "F"): wp.column_dimensions[col].width = 17
wp.column_dimensions["G"].width = 10
wp.freeze_panes = "A5"
wp.sheet_view.showGridLines = False

# ---------------------------------------------------------------- hojas originales
# Ninguna hoja se copia del libro original: el libro se reconstruye entero aqui.
# (la hoja «Fondo de emergencia» se construye entera mas abajo)

# ---------------------------------------------------------------- graficos
def quesito(ws_dst, titulo_g, labels_ref, data_ref, anchor):
    ch = PieChart(); ch.title = titulo_g; ch.height = 8; ch.width = 13
    ch.add_data(data_ref, titles_from_data=False); ch.set_categories(labels_ref)
    ch.dataLabels = openpyxl.chart.label.DataLabelList(); ch.dataLabels.showPercent = True
    ws_dst.add_chart(ch, anchor)

filas_bloque = [b[0] for b in bloques]
wg = wb.create_sheet("Gráficos")
wg.cell(1, 1, "Bloque"); wg.cell(1, 2, "Anual")
for i, (fb, nom, _s) in enumerate(bloques, start=2):
    wg.cell(i, 1, nom); wg.cell(i, 2, f"={IG}!D{fb}").number_format = EUR
n = len(bloques) + 1
quesito(wg, "Gastos por bloque",
        Reference(wg, min_col=1, min_row=2, max_row=n), Reference(wg, min_col=2, min_row=2, max_row=n), "D2")
wg.cell(n + 3, 1, "Ingresos"); wg.cell(n + 3, 2, "Anual")
for i, (fg, nom, _f, _l) in enumerate(grupos_ing, start=n + 4):
    wg.cell(i, 1, nom); wg.cell(i, 2, f"={IG}!D{fg}").number_format = EUR
quesito(wg, "Ingresos por origen",
        Reference(wg, min_col=1, min_row=n + 4, max_row=n + 3 + len(grupos_ing)),
        Reference(wg, min_col=2, min_row=n + 4, max_row=n + 3 + len(grupos_ing)), "D20")
wg.column_dimensions["A"].width = 24; wg.column_dimensions["B"].width = 14
wg.sheet_view.showGridLines = False


# ---------------------------------------------------------------- PATRIMONIO NETO
import datetime
wn = wb.create_sheet("Patrimonio Neto")
titulo(wn, 1, "PATRIMONIO NETO", "Lo que tienes menos lo que debes. CIFRAS INVENTADAS: pon las tuyas y actualiza la fecha de E1 (los plazos de «Objetivos» se cuentan desde ese año).")
wn["D1"] = "Última actualización"; wn["D1"].font = Font(bold=True, size=9, color="808080")
wn["E1"] = datetime.date(2026, 9, 1); wn["E1"].number_format = "dd/mm/yyyy"
wn["E1"].font = Font(bold=True, size=9, color="808080")

wn["A3"] = "PATRIMONIO NETO (Activos − Pasivos)"
wn["A3"].font = Font(bold=True, size=14, color="006100")
wn["E3"] = "=E5-E18"; wn["E3"].number_format = EUR
wn["E3"].font = Font(bold=True, size=14, color="006100")
for j in range(1, 6): wn.cell(3, j).fill = PatternFill("solid", fgColor=VERDE)

ACTIVOS = [
 ("Vivienda habitual", "", "", "Ana 50% / Luis 50%", 260000.00, "EJEMPLO. Valor estimado de mercado hoy (no el de compra). La hipoteca va abajo, en pasivos"),
 ("Cuenta corriente común", "Banco A", "*1234", "Común", 4500.00, "EJEMPLO. La del día a día: aquí entran las nóminas y salen los recibos"),
 ("Cuenta remunerada (colchón)", "Banco B", "", "Común", 14000.00, "EJEMPLO. Fondo de emergencia, separado de la cuenta del día a día para no tocarlo"),
 ("Cuenta personal Ana", "", "", "Ana", 1500.00, ""),
 ("Cuenta personal Luis", "", "", "Luis", 1200.00, ""),
 ("Cartera conservadora", "Gestora X", "", "Común", 3000.00, "EJEMPLO. Dinero para objetivos a menos de ~10 años. La pestaña «Objetivos» lee de aquí lo ya invertido"),
 ("Cartera indexada global", "Gestora Y", "", "Común", 18000.00, "EJEMPLO. Dinero para objetivos lejanos (15+ años). La pestaña «Objetivos» lee de aquí lo ya invertido"),
 ("Plan de pensiones", "Gestora Z", "", "Luis", 9000.00, "EJEMPLO. No es dinero disponible: solo se rescata en los supuestos que marca la ley"),
 (None, None, None, None, None, None),
]
PASIVOS = [
 ("Hipoteca de la vivienda", "Banco A", "", "Ana 50% / Luis 50%", 165000.00,
  "EJEMPLO. Capital pendiente (lo que aún debes), no la cuota"),
 (None, None, None, None, None, None), (None, None, None, None, None, None),
]
FILA_ACT = {it[0]: 7 + i for i, it in enumerate(ACTIVOS) if it[0]}
PN = lambda nombre: f"'Patrimonio Neto'!E{FILA_ACT[nombre]}"
def tabla(wn, r0, titulo_b, items, color):
    wn.cell(r0, 1, titulo_b).font = Font(bold=True, size=12, color="FFFFFF")
    for j in range(1, 6): wn.cell(r0, j).fill = PatternFill("solid", fgColor=color)
    wn.cell(r0, 5, f"=SUM(E{r0+2}:E{r0+1+len(items)})").number_format = EUR
    wn.cell(r0, 5).font = Font(bold=True, size=12, color="FFFFFF")
    for j, h in enumerate(["Descripción", "Entidad", "Nº de cuenta", "Titular/es", "Valoración"], start=1):
        c = wn.cell(r0+1, j, h); c.font = Font(bold=True, size=9, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="8EA9DB")
    for i, it in enumerate(items):
        rr = r0 + 2 + i
        for j in range(4):
            if it[j] is not None: wn.cell(rr, j+1, it[j])
            wn.cell(rr, j+1).font = Font(size=10)
        if it[4] is not None: wn.cell(rr, 5, it[4])
        wn.cell(rr, 5).number_format = EUR0
        if it[5]:
            n = wn.cell(rr, 7, it[5]); n.font = Font(size=8, color="808080")
        for j in range(1, 6): wn.cell(rr, j).border = BORDE
    return r0 + 2 + len(items)

fin_a = tabla(wn, 5, "ACTIVOS (BIENES)", ACTIVOS, "4472C4")
r_pas = fin_a + 1
fin_p = tabla(wn, r_pas, "PASIVOS (DEUDAS)", PASIVOS, "C00000")
wn["E3"] = f"=E5-E{r_pas}"
wn.cell(r_pas + len(PASIVOS) + 3, 1, "RATIO DE ENDEUDAMIENTO").font = Font(bold=True, size=11, color=AZUL)
wn.cell(r_pas + len(PASIVOS) + 3, 5, f"=IFERROR(E{r_pas}/E5,0)").number_format = PCT
wn.cell(r_pas + len(PASIVOS) + 3, 5).font = Font(bold=True, size=11, color=AZUL)
wn.cell(r_pas + len(PASIVOS) + 4, 1, "Deudas sobre activos. Por debajo del 40% se considera cómodo.").font = Font(size=8, italic=True, color="808080")
wn.column_dimensions["A"].width = 36
for col in ("B", "C", "D"): wn.column_dimensions[col].width = 18
wn.column_dimensions["E"].width = 16
wn.column_dimensions["F"].width = 3
wn.column_dimensions["G"].width = 95
wn.sheet_view.showGridLines = False
wb.move_sheet("Patrimonio Neto", offset=-(len(wb.sheetnames)-1))


# ---------------------------------------------------------------- FONDO DE EMERGENCIA
wf = wb.create_sheet("Fondo de emergencia")
titulo(wf, 1, "FONDO DE EMERGENCIA",
       "Cuánto aguantáis sin ingresos. La clave: en una emergencia NO se gasta lo mismo — los revisables se cortan el primer día.")
def bloq(wf, r, txt):
    c = wf.cell(r, 1, txt); c.font = Font(bold=True, size=12, color="FFFFFF")
    for j in range(1, 6): wf.cell(r, j).fill = PatternFill("solid", fgColor="4472C4")
    wf.row_dimensions[r].height = 18
def lin(wf, r, txt, formula, nota=None, negrita=False, fmt=EUR, ind=1):
    c = wf.cell(r, 1, ("   " * ind) + txt)
    c.font = Font(bold=negrita, size=11 if negrita else 10, color=AZUL if negrita else "000000")
    cc = wf.cell(r, 3, formula); cc.number_format = fmt
    cc.font = Font(bold=negrita, size=11 if negrita else 10, color=AZUL if negrita else "000000")
    if nota:
        n = wf.cell(r, 5, nota); n.font = Font(size=8, color="808080")
    for j in range(1, 4): wf.cell(r, j).border = BORDE

bloq(wf, 3, "A) ¿CUÁNTO CUESTA VIVIR EN MODO EMERGENCIA?")
lin(wf, 4, "Gasto actual completo", f"={IG}!D{r_tot_gas}/12", "Todo lo de la hoja de gastos", ind=1)
lin(wf, 5, "− Revisables (se cortan el primer día)", f"=-{IG}!D{bloques[2][0]}/12", "Suscripciones, ocio, vacaciones, restaurantes, imprevistos, compras online", ind=1)
lin(wf, 6, "− Plan de pensiones", f"=-{IG}!D{FILA['Jubilación (plan de pensiones)']}/12", "Se puede suspender en una emergencia", ind=1)
lin(wf, 7, "− Supermercado más ajustado", -80.00, "EJEMPLO de ajuste — cámbialo o bórralo", ind=1)
for _i, _r in enumerate(range(8, 12)):
    lin(wf, _r, "− (escribe aquí el concepto)", None,
        "LIBRE: pon el concepto a la izquierda y en negativo lo que recortarías al mes. Ej.: supermercado −150, ocio −80, ropa −40" if _i == 0 else None, ind=1)
    wf.cell(_r, 1).font = Font(size=10, italic=True, color="A6A6A6")
lin(wf, 12, "GASTO DE SUPERVIVENCIA (al mes)", "=SUM(C4:C11)", "Es la cifra que manda en todo lo de abajo", negrita=True, ind=0)
wf.cell(12, 3).fill = PatternFill("solid", fgColor="FFF2CC")

bloq(wf, 14, "B) INGRESOS DISPONIBLES")
lin(wf, 15, "Neto mensual Ana", f"={IG}!B{FILA['Nómina neta Ana']}", "Solo la nómina: lo que se paga en especie y los extras variables se pierden con el empleo", ind=1)
lin(wf, 16, "Neto mensual Luis", f"={IG}!B{FILA['Nómina neta Luis']}", "Mismo criterio", ind=1)
lin(wf, 17, "Paro neto estimado Ana", 1150.00,
    "EJEMPLO. Calcula el tuyo en la web del SEPE: depende de tu base de cotización, tiene un tope que sube con los hijos a cargo y baja a partir del día 181", ind=1)
lin(wf, 18, "Paro neto estimado Luis", 1050.00,
    "EJEMPLO. Mismo criterio. La DURACIÓN se pone en el bloque G", ind=1)

bloq(wf, 20, "C) ESCENARIOS")
for j2, h in enumerate(["Escenario", "Ingresos/mes", "Déficit/mes", "Meses con el colchón", "Meses con toda la liquidez"], start=1):
    c = wf.cell(21, j2, h); c.font = Font(bold=True, size=9, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="8EA9DB"); c.alignment = Alignment(horizontal="center", wrap_text=True)
wf.row_dimensions[21].height = 30
esc = [("Ana en paro, Luis trabaja", "=C16+C17"),
       ("Luis en paro, Ana trabaja", "=C15+C18"),
       ("Los dos en paro (cobrando prestación)", "=C17+C18"),
       ("Los dos en paro, PRESTACIÓN AGOTADA", 0)]
for i2, (nom, fing) in enumerate(esc):
    r = 22 + i2
    wf.cell(r, 1, nom).font = Font(bold=(i2 == 3), size=10, color="C00000" if i2 == 3 else "000000")
    wf.cell(r, 2, fing).number_format = EUR
    wf.cell(r, 3, f"=MAX(0,$C$12-B{r})").number_format = EUR
    wf.cell(r, 4, f"=IF(C{r}>0,$C$33/C{r},0)").number_format = '0.0'
    wf.cell(r, 5, f"=IF(C{r}>0,$C$37/C{r},0)").number_format = '0.0'
    for j2 in range(1, 6): wf.cell(r, j2).border = BORDE
wf.cell(26, 1, "Un 0 en «meses» significa que NO hay déficit: ese sueldo cubre el gasto de supervivencia sin tocar el colchón. La última fila es el escenario a vigilar: cuando se acaba la prestación.").font = Font(size=8, italic=True, color="808080")

bloq(wf, 28, "D) LIQUIDEZ DISPONIBLE")
lin(wf, 29, "Colchón de emergencia", "=" + PN("Cuenta remunerada (colchón)"), "Intocable, designado para esto", ind=1)
lin(wf, 30, "Cuenta corriente común", "=" + PN("Cuenta corriente común"), "Buffer operativo + lo pendiente de invertir", ind=1)
lin(wf, 31, "Cuentas personales", "=" + PN("Cuenta personal Ana") + "+" + PN("Cuenta personal Luis"), None, ind=1)
lin(wf, 33, "COLCHÓN DESIGNADO", "=C29", "El que se usa en la columna «meses con el colchón»", negrita=True, ind=0)
lin(wf, 37, "LIQUIDEZ TOTAL", "=C29+C30+C31", "Todo el dinero disponible sin vender inversiones. NO incluye las carteras de inversión ni el plan de pensiones: vender inversiones en mitad de una crisis es justo lo que el colchón evita", negrita=True, ind=0)
for rr in (33, 37):
    for j2 in range(1, 4): wf.cell(rr, j2).fill = PatternFill("solid", fgColor=VERDE)

bloq(wf, 39, "E) OBJETIVO DEL FONDO — elige tú el nivel")
for j2, h in enumerate(["Escenario", "Déficit/mes", "Cubrir 6 meses", "Cubrir 12 meses", "Sobra/falta frente al colchón (12 m)"], start=1):
    c = wf.cell(40, j2, h); c.font = Font(bold=True, size=9, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="8EA9DB"); c.alignment = Alignment(horizontal="center", wrap_text=True)
wf.row_dimensions[40].height = 30
for i2, (nom, orig) in enumerate([("Ana en paro, Luis trabaja", 22),
                                  ("Luis en paro, Ana trabaja", 23),
                                  ("Los dos en paro (cobrando prestación)", 24),
                                  ("Los dos en paro, PRESTACIÓN AGOTADA", 25)]):
    r = 41 + i2
    wf.cell(r, 1, nom).font = Font(size=10, color="C00000" if i2 == 3 else "000000")
    wf.cell(r, 2, f"=C{orig}").number_format = EUR
    wf.cell(r, 3, f"=C{orig}*6").number_format = EUR
    wf.cell(r, 4, f"=C{orig}*12").number_format = EUR
    wf.cell(r, 5, f"=$C$33-D{r}").number_format = EUR
    for j2 in range(1, 6): wf.cell(r, j2).border = BORDE
r = 45
wf.cell(r, 1, "MEDIA de 3 casos (los de déficit 0 cuentan como uno)").font = Font(bold=True, size=10, color=AZUL)
wf.cell(r, 2, "=(MAX(C22,C23)+C24+C25)/3").number_format = EUR
wf.cell(r, 3, "=B45*6").number_format = EUR
wf.cell(r, 4, "=B45*12").number_format = EUR
wf.cell(r, 5, "=$C$33-D45").number_format = EUR
for j2 in range(1, 6):
    wf.cell(r, j2).fill = PatternFill("solid", fgColor=GRIS); wf.cell(r, j2).border = BORDE
    wf.cell(r, j2).font = Font(bold=True, size=10, color=AZUL)
wf.cell(46, 1, "«Ana en paro» y «Luis en paro» son el MISMO suceso (cae uno de los dos), así que cuentan como un solo caso, con el peor de los dos (MAX). Aun así la media da el mismo peso a «cae uno» que a «caen los dos y se agota la prestación», que es mucho menos probable. Úsala como referencia alta, no como objetivo.").font = Font(size=8, italic=True, color="808080")

bloq(wf, 48, "F) LO QUE EL COLCHÓN CUBRE ADEMÁS DEL PARO")
lin(wf, 49, "Desfase hasta el primer cobro del paro", "=C4*2",
    "2 meses de gasto COMPLETO (no de supervivencia): no cobras el día 1, hay que solicitarlo y el primer pago tarda semanas. Cambia el 2 si quieres otro margen", ind=1)
lin(wf, 50, "Imprevisto gordo (avería, derrama, salud)", 1500.00,
    "EJEMPLO. Es el uso MÁS FRECUENTE del colchón: caldera, derrama de la comunidad, avería seria del coche, gasto médico. Pon tu cifra", ind=1)
lin(wf, 51, "SUELO DEL COLCHÓN (sin contar el paro)", "=C49+C50", "Esto haría falta aunque nadie perdiera el empleo", negrita=True, ind=0)
for j2 in range(1, 4): wf.cell(51, j2).fill = PatternFill("solid", fgColor="FFF2CC")
lin(wf, 53, "Colchón actual", "=C33", None, negrita=True, ind=0)
for j2 in range(1, 4): wf.cell(53, j2).fill = PatternFill("solid", fgColor=VERDE)
wf.cell(55, 1, "Cómo leer esto: si cae uno solo de los dos y el otro sueldo cubre el gasto de supervivencia, el déficit es CERO. El trabajo real del colchón suele ser el bloque F: el hueco hasta que entra el paro y el imprevisto gordo.").font = Font(size=9, italic=True, color="404040")
wf.cell(56, 1, "Si pierdes el empleo también pierdes lo que la empresa paga en especie (p. ej. el seguro médico). Si sigues contándolo como gasto de supervivencia, el cálculo va por el lado seguro.").font = Font(size=8, italic=True, color="808080")


bloq(wf, 58, "G) SI CAÉIS LOS DOS: ¿CUÁNTO AGUANTÁIS DE VERDAD?")
lin(wf, 59, "Meses que dura la prestación", 24,
    "EJEMPLO. Depende de lo cotizado en los últimos 6 años; el máximo son 720 días (24 meses). Es el dato que manda en todo este bloque", ind=1)
lin(wf, 60, "Déficit mientras cobráis la prestación", "=C24", "Del escenario «los dos en paro, cobrando»", ind=1)
lin(wf, 61, "Déficit cuando se agota", "=C25", "Del escenario «prestación agotada»", ind=1)
lin(wf, 62, "Coste de los meses con prestación", "=C59*C60", "Lo que os cuesta el agujero durante toda la prestación", ind=1)
lin(wf, 63, "Queda del colchón después", "=C33-C62", None, ind=1)
lin(wf, 64, "Meses extra que aguanta el colchón", "=MAX(0,C63/C61)", None, fmt='0.0', ind=1)
lin(wf, 65, "MESES TOTALES con el colchón", "=C59+C64", "Prestación + lo que estira el colchón después", negrita=True, fmt='0.0', ind=0)
for j2 in range(1, 4): wf.cell(65, j2).fill = PatternFill("solid", fgColor="FFF2CC")
lin(wf, 67, "Queda de la liquidez total después", "=C37-C62", None, ind=1)
lin(wf, 68, "Meses extra con toda la liquidez", "=MAX(0,C67/C61)", None, fmt='0.0', ind=1)
lin(wf, 69, "MESES TOTALES con toda la liquidez", "=C59+C68", "Sin vender ninguna inversión", negrita=True, fmt='0.0', ind=0)
for j2 in range(1, 4): wf.cell(69, j2).fill = PatternFill("solid", fgColor=VERDE)
wf.cell(71, 1, "Este bloque es el que responde de verdad a «¿cuánto aguantamos?». Los meses de prestación son casi gratis (el déficit es pequeño); el reloj empieza a correr el día que se agota.").font = Font(size=9, italic=True, color="404040")

wf.column_dimensions["A"].width = 44
wf.column_dimensions["B"].width = 16
wf.column_dimensions["C"].width = 16
wf.column_dimensions["D"].width = 18
wf.column_dimensions["E"].width = 95
wf.sheet_view.showGridLines = False

import sys
AQUI = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(AQUI, "nuevo.xlsx")
wb.save(OUT)
print("guardado:", OUT)
json.dump({"pp_anual": f"'Ingresos - Gastos'!D{FILA['Jubilación (plan de pensiones)']}", "ahorro": f"'Ingresos - Gastos'!D{R_AHORRO + 1}",
           "inv": {k: PN(k) for k in ("Cartera conservadora", "Cartera indexada global", "Plan de pensiones")}},
          open(os.path.join(AQUI, "refs.json"), "w"), ensure_ascii=False)
print("filas clave -> total ingresos:", r_tot_ing, "| total gastos:", r_tot_gas, "| ahorro:", R_AHORRO, "| presupuesto total:", R_PRES_TOT)
